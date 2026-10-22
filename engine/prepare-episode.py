"""Prepare an episode folder (any channel) and build its edit in one go.

  1. read the script spreadsheet (sheet "Content") into script.json
  2. map every script section to its VO files in EN/AUDIO
  3. transcribe the VO with whisper.cpp (word timestamps) so lines can be timed
  4. write episode.json and run build-timeline.py (Resolve .fcpxml + Premiere .xml)

Usage:
  python prepare-episode.py "<episode folder>" [--sections all|intro,1,quiz,2,3,outro]
                            [--name NAME] [--out DIR] [--no-build]

Episode folder layout (as used by the channels):
  <root>/*.xlsx (anywhere, sheet "Content"), <root>/VIDEO/<n>-<ID>.*,
  VO in <root>/EN/AUDIO EN (Tripinsight; the edit is then written to <root>/EN), <root>/EN/AUDIO or <root>/AUDIO:
  {INTRO .. 1|2, STORY .. 1|2|3, QUIZ .., 1. <Place> EN 1|2|3, OUTRO ..}.mp3,
  episode music in <root>/Epidemic Sound/Music or <root>/SOUND/Music (optional)
"""
import argparse, glob, json, os, re, subprocess, sys

HERE = os.path.dirname(os.path.abspath(__file__))
MODEL = os.path.join(HERE, "models", "ggml-base.en.bin")
sys.path.insert(0, HERE)
import channel

ap = argparse.ArgumentParser()
ap.add_argument("root")
ap.add_argument("--sections", default="all", help="all, or e.g. intro,1,quiz,2,3,outro")
ap.add_argument("--name")
ap.add_argument("--region", help="region shown in lower thirds / opening title (default: folder name after ' - ')")
ap.add_argument("--out")
ap.add_argument("--template", help="default: the channel's template")
ap.add_argument("--music-lib", help="default: the channel's music library")
ap.add_argument("--no-build", action="store_true")
args = ap.parse_args()
CH = channel.load(args.root)        # channels/<slug>/channel.json, picked from the episode folder
args.template = args.template or CH["template"]
args.music_lib = args.music_lib or CH["music_lib"]
root = os.path.abspath(args.root.rstrip("/"))
name = args.name or os.path.basename(root)
region = args.region or (os.path.basename(root).split(" - ")[-1] if " - " in os.path.basename(root) else "")
# VO: <root>/EN/AUDIO EN (Tripinsight: one folder per language, each with its own project),
# <root>/EN/AUDIO (Earthcue) or <root>/AUDIO (Docs Untold)
audio_dir = next((d for d in (os.path.join(root, "EN", "AUDIO EN"), os.path.join(root, "EN", "AUDIO"),
                              os.path.join(root, "AUDIO")) if os.path.isdir(d)), None)
if not audio_dir:
    sys.exit("no EN/AUDIO folder")
lang_dir = os.path.dirname(audio_dir) if os.path.basename(audio_dir) == "AUDIO EN" else root
out_dir = args.out or lang_dir                # <episode>/<name>.xml, MEDIA (music), FINAL, .autoedit
work = os.path.join(out_dir, ".autoedit", "_work")
os.makedirs(os.path.join(work, "vo"), exist_ok=True)


def visible(paths):
    return sorted(p for p in paths if not os.path.basename(p).startswith("._"))


# ---------------------------------------------------------------- 1. script
import openpyxl

books = visible(glob.glob(os.path.join(root, "**", "*.xlsx"), recursive=True))
book = next((b for b in books if "Content" in openpyxl.load_workbook(b, read_only=True).sheetnames), None)
if not book:
    sys.exit(f"no .xlsx with a 'Content' sheet under {root}")
ws = openpyxl.load_workbook(book, data_only=True)["Content"]
header = [str(c or "").strip() for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
col = {h: i for i, h in enumerate(header)}
c_num, c_en = col.get("Number", 1), col.get("English", 2)
c_file = col.get("File name EN", 10)
def is_sentence(text):
    """A full spoken sentence (not a heading) that lost its Number: one line, 5+ words, ends with . ! ?"""
    return "\n" not in text and len(text.split()) >= 5 and re.search(r"[.?!][\"”’')]*$", text)


rows, sections, sec = [], [], None
for r in ws.iter_rows(min_row=2, values_only=True):
    n, en = r[c_num], r[c_en]
    if en is None and n is None:
        continue
    f = str(r[c_file] or "").strip()
    if n is None and en and not f:
        text = str(en).strip()
        if is_sentence(text) and rows and rows[-1]["sec"] == sec:
            rows[-1]["en"] = rows[-1]["en"].rstrip() + " " + text   # spoken over the previous line's footage
            print(f"note: line without Number/footage joined to line {rows[-1]['n']}: {text[:60]}")
            continue
        if (sec and "\n" in text and not re.search(r"\nA[.)]", text)
                and not any(x["sec"] == sec for x in rows)):
            continue                                     # the whole section's text pasted under its heading
        # a section heading row; Chaseland headings start with a timestamp ("15:19 Sydney Harbour"): dropped
        sec = re.sub(r"^\d{1,2}:\d{2}(?::\d{2})?\s+", "", text)
        sections.append(sec)
        continue
    if sec and f:                                        # a line with footage (cold-open lines may lack a Number)
        text = "" if en is None else str(en)
        # a line with footage but no text shares the previous line's time on screen
        rows.append(dict(n=n, sec=sec, en=text, file=f, **({"extra": True} if not text.strip() else {})))


def kind_of(i, title):
    t = title.lower()
    if re.match(r"\d+\.", title):
        return "location"
    if i == 0 or t.startswith("intro"):
        return "intro"
    if re.search(r"\nA[.)]", title) or "question" in t or "quiz" in t:
        return "quiz"
    if re.match(r"out+ro", t):
        return "outro"
    return "story"


# ---------------------------------------------------------------- 2. VO files
audio =[f for f in visible(os.listdir(audio_dir)) if f.lower().endswith((".mp3", ".wav"))]


def fold(s):
    return re.sub(r"\s+", " ", s).strip().lower()


def named(f, title):
    """VO file f is named after the section title (its number / "0." prefix and part number aside)."""
    return re.match(rf"{re.escape(fold(title))}(\s|\d|\.|$)", fold(re.sub(r"^\d+\.\s*", "", f))) is not None


# a location heading without its number (Chaseland: "Sydney Harbour") takes the number of its VO files
# ("1. Sydney Harbour 1.mp3"), so it becomes "1. Sydney Harbour" like every other channel's headings
for i, title in enumerate(sections):
    if i == 0 or "\n" in title or re.match(r"\d+\.", title):
        continue
    num = next((re.match(r"(\d+)\.", f).group(1) for f in audio
                if re.match(r"[1-9]\d*\.", f) and named(f, title)), None)
    if num:
        sections[i] = f"{num}. {title}"
        for r in rows:
            if r["sec"] == title:
                r["sec"] = sections[i]
json.dump(rows, open(os.path.join(work, "script.json"), "w"), ensure_ascii=False, indent=0)


def part_no(f):
    """Part number of a VO file: '... EN 2.mp3', '... EN 2 - Patagonia.mp3' or '... 2.mp3'."""
    m = re.search(r"\bEN\s*(\d+)", f) or re.search(r"(\d+)\s*\.\w+$", f)
    return int(m.group(1)) if m else 0


def vo_for(kind, title):
    if kind == "location":
        num = re.match(r"(\d+)\.", title).group(1)
        files = [f for f in audio if re.match(rf"{num}\.(?!\d)", f)]   # "5. Ushuaia" and "5.Ushuaia"
    else:
        key = {"intro": "INTRO", "quiz": "QUIZ", "outro": "OUTRO", "story": "STORY"}[kind]
        files = ([f for f in audio if f.upper().startswith(key)]
                 or [f for f in audio if not re.match(r"[1-9]\d*\.", f) and named(f, title)])  # "History 1.mp3"
    return sorted(files, key=part_no)


plan = []
for i, title in enumerate(sections):
    k = kind_of(i, title)
    plan.append(dict(prefix=title, kind=k, files=vo_for(k, title),
                     key=re.match(r"(\d+)\.", title).group(1) if k == "location" else k))


def words(s):
    return re.sub(r"[^a-z0-9 ]", "", s.lower().replace("_", " ")).split()


# a VO file named after its own text (e.g. "_... Welcome to Arka.mp3" for a cold-open line) goes to the
# section whose first line it starts; part_no 0 keeps it ahead of the numbered parts
claimed = {f for p in plan for f in p["files"]}
title_vo = None
for f in [f for f in audio if f not in claimed]:
    if region and fold(os.path.splitext(f)[0]) == fold(region):
        title_vo = f                                     # the region name said over the opening title (Chaseland)
        print(f"note: VO {f} is the opening title's voice")
        continue
    fw = words(os.path.splitext(f)[0])
    for p in plan:
        first = next((r["en"] for r in rows if r["sec"] == p["prefix"]), "")
        lw = words(first)
        if fw and len(lw) >= len(fw) and lw[:len(fw) - 1] == fw[:-1] and lw[len(fw) - 1].startswith(fw[-1]):
            p["files"] = sorted(p["files"] + [f], key=part_no)
            print(f"note: VO {f} -> {p['prefix'].split(chr(10))[0]} (starts its first line)")
            break
    else:
        print(f"warning: VO file not used: {f}")
if args.sections != "all":
    want = [w.strip().lower() for w in args.sections.split(",")]
    plan = [next(p for p in plan if p["key"] == w) for w in want]
missing = [p["prefix"].split("\n")[0] for p in plan if not p["files"]]
if missing:
    sys.exit("no VO files found for: " + ", ".join(missing))

# ---------------------------------------------------------------- 3. transcription
index, n = [], 0
for p in plan:
    p["vo"] = []
    for f in p["files"]:
        n += 1
        p["vo"].append(n)
        index.append(f"{n:02d}\t{f}")
        wav, csv_base = os.path.join(work, "vo", f"{n:02d}.wav"), os.path.join(work, "vo", f"{n:02d}")
        if not os.path.exists(csv_base + ".csv"):
            subprocess.run(["ffmpeg", "-v", "error", "-y", "-i", os.path.join(audio_dir, f), "-ar", "16000",
                            "-ac", "1", wav], check=True)
            subprocess.run(["whisper-cli", "-m", MODEL, "-f", wav, "-ml", "1", "-sow", "-ocsv", "-of", csv_base,
                            "-np"], check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        print(f"VO {n:02d}  {f}")
open(os.path.join(work, "vo", "index.tsv"), "w").write("\n".join(index) + "\n")

# ---------------------------------------------------------------- 4. episode.json + build
music_dir = next((d for d in (os.path.join(root, "Epidemic Sound", "Music"), os.path.join(root, "SOUND", "Music"),
                              os.path.join(root, "PROJECT", "Epidemic Sound", "Music"))
                  if os.path.isdir(d)), None)
episode = dict(channel=CH["slug"], root=root, name=name, region=region, out_dir=out_dir, audio_dir=audio_dir, template=args.template,
               music_lib=args.music_lib,
               music=[os.path.join(music_dir, f) for f in visible(os.listdir(music_dir))
                      if f.lower().endswith((".mp3", ".wav"))] if music_dir else [],
               title_vo=title_vo,
               sections=[dict(prefix=p["prefix"], kind=p["kind"], vo=p["vo"]) for p in plan])
json.dump(episode, open(os.path.join(work, "episode.json"), "w"), ensure_ascii=False, indent=1)
print("sections:", ", ".join(f'{p["prefix"].split(chr(10))[0]} ({p["kind"]}, {len(p["vo"])} VO)' for p in plan))
if not args.no_build:
    subprocess.run([sys.executable, "-I", os.path.join(HERE, "build-timeline.py"), work], check=True)
    print("\nDone ->", os.path.join(out_dir, name + ".xml"))
